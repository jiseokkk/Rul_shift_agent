"""여러 run 을 시나리오 단위 v2 규약(docs/eval_v2_scenario.md)으로 채점해 나란히 비교한다.

  python agent/scripts/compare_scenario_runs.py --runs <run_id> ... --labels <name> ... --out or500 [--ref deepseek-v4-pro]

각 run 은 evaluate_run (→ runs/{id}/eval/) 으로 채점하고, 결과를 reports/eval_v2/ 에 모은다:
  compare_{out}.md / .xlsx   메인 표 (기준선 포함) · unit bootstrap CI · 유형별 DR · 센서별 · 시점별 오염 전 경보율 · --ref 대비 짝지은 차이 CI
  {out}_{label}.md           run 별 보고서 사본
"""
from __future__ import annotations

import argparse
import shutil

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from _bootstrap import AGENT, load_cfg
from src.eval.report import CI_METRICS, evaluate_run
from src.eval.scenario_table import MAIN_COLS, metrics_by, pre_contam_by_timing
from src.eval.stats import paired_diff_bootstrap

OUT = AGENT / "reports" / "eval_v2"
NAMES = {"n_degraded": "저하 n", "detection_rate": "DR", "pre_contam_rate": "PreContam", "pre_degr_rate": "PreDegr", "miss_rate": "Miss",
         "delay_median": "Delay 중앙값", "n_nondegraded": "비저하 n", "scenario_FAR": "FAR", "FAR_clean": "FAR clean", "FAR_contam": "FAR contam",
         "isolation_rate": "Isolation"}


def md(df: pd.DataFrame, fmt: str = "{:.3f}") -> str:
    cols = list(df.columns)
    L = ["| " + " | ".join(map(str, cols)) + " |", "|" + "---|" * len(cols)]
    for _, r in df.iterrows():
        cells = []
        for c in cols:
            v = r[c]
            if isinstance(v, float) and not pd.isna(v):
                cells.append(f"{int(v)}" if float(v).is_integer() else fmt.format(v))
            else:
                cells.append("" if pd.isna(v) else str(v))
        L.append("| " + " | ".join(cells) + " |")
    return "\n".join(L)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--runs", nargs="+", required=True)
    ap.add_argument("--labels", nargs="*", default=None)
    ap.add_argument("--out", default="or500")
    ap.add_argument("--ref", default=None, help="짝지은 차이의 기준 label (기본: 첫 번째)")
    ap.add_argument("--n-boot", type=int, default=None)
    a = ap.parse_args()
    acfg = load_cfg("agent")
    labels = a.labels or a.runs
    ref = a.ref or labels[0]
    OUT.mkdir(parents=True, exist_ok=True)

    res, S = {}, {}
    for rid, lab in zip(a.runs, labels):
        rd = AGENT / "runs" / rid
        res[lab] = evaluate_run(rd, acfg, None, a.n_boot)
        S[lab] = pd.read_csv(rd / "eval" / "scenario_table.csv")
        shutil.copy(rd / "eval" / "report.md", OUT / f"{a.out}_{lab}.md")
        print(f"{lab}: DR {res[lab]['metrics']['detection_rate']:.3f}  FAR {res[lab]['metrics']['scenario_FAR']:.3f}")

    e = res[ref]["config"]
    main_rows = [{"method": r["method"], **{NAMES[c]: r[c] for c in MAIN_COLS}} for r in res[ref]["baselines"]
                 if not r["method"].startswith("무작위 p=0.0")]  # 공통 기준선 (항상 0/1, p=0.05) 은 기준 run 에서 한 번
    main_rows.append({"method": "무작위 p=0.050", **{NAMES[c]: r[c] for c in MAIN_COLS for r in res[ref]["baselines"] if r["method"] == "무작위 p=0.050"}})
    for lab in labels:
        pm = [r for r in res[lab]["baselines"] if "에이전트 경보율" in r["method"]][0]
        main_rows.append({"method": f"무작위 (p={res[lab]['config']['p_match']:.3f}, {lab} 경보율)", **{NAMES[c]: pm[c] for c in MAIN_COLS}})
        main_rows.append({"method": f"**{lab}**", **{NAMES[c]: res[lab]["metrics"][c] for c in MAIN_COLS}})
    main = pd.DataFrame(main_rows)

    ci_rows = []
    for lab in labels:
        for r in res[lab]["ci"] or []:
            ci_rows.append({"metric": NAMES.get(r["metric"], r["metric"]), "run": lab, "point": r["point"], "95% CI": f"[{r['lo']:.3f}, {r['hi']:.3f}]"})
    ci = pd.DataFrame(ci_rows)

    by_type = pd.concat([metrics_by(S[lab], ["type"], ["n_degraded", "detection_rate", "pre_contam_rate", "pre_degr_rate", "miss_rate", "scenario_FAR"]).assign(run=lab)
                         for lab in labels]).pivot(index="type", columns="run", values=["detection_rate", "scenario_FAR"]).round(3)
    by_type.columns = [f"{NAMES[m]} {r}" for m, r in by_type.columns]
    by_type = by_type.reset_index()
    by_sensor = pd.concat([metrics_by(S[lab], ["sensor"], ["detection_rate", "scenario_FAR", "isolation_rate"]).assign(run=lab) for lab in labels]) \
        .pivot(index="sensor", columns="run", values=["detection_rate", "scenario_FAR"]).round(3)
    by_sensor.columns = [f"{NAMES[m]} {r}" for m, r in by_sensor.columns]
    by_sensor = by_sensor.reset_index()
    timing = pd.concat([pre_contam_by_timing(S[lab]).assign(run=lab) for lab in labels]) \
        .pivot(index="timing_p", columns="run", values="pre_contam_rate").round(3)
    timing.insert(0, "오염 전 길이 중앙값", pre_contam_by_timing(S[ref]).set_index("timing_p")["pre_len_median"])
    timing = timing.reset_index()

    diff_rows = []
    for lab in labels:
        if lab == ref:
            continue
        D = paired_diff_bootstrap(S[lab], S[ref], CI_METRICS, n_boot=int(a.n_boot if a.n_boot is not None else acfg["eval"]["n_boot"]))
        for r in D.itertuples():
            diff_rows.append({"비교": f"{lab} − {ref}", "metric": NAMES.get(r.metric, r.metric), "diff": r.diff, "95% CI": f"[{r.lo:.3f}, {r.hi:.3f}]",
                              "유의": "*" if (r.lo > 0 or r.hi < 0) else ""})
    diff = pd.DataFrame(diff_rows)

    n = res[ref]["metrics"]
    L = [f"# run 비교 (시나리오 단위 v2) — {a.out}", "",
         f"작성 {pd.Timestamp.now():%Y-%m-%d}. 규약 `docs/eval_v2_scenario.md`: 첫 경보 원칙 (cycle {e['judge_from']} 이후 첫 경보, 오염 전 포함), "
         f"Δ = {e['Delta']}, w = {e['w']}, H = {e['H']}, k = {e['k']}. 저하 {n['n_degraded']} / 비저하 {n['n_nondegraded']} 시나리오, unit {S[ref]['unit'].nunique()}. "
         "LLM 재호출 없이 기존 판정 로그를 재채점했다 (v1 실행은 수명 끝까지 판정했으므로 창이 모두 덮인다).", "",
         "run: " + " · ".join(f"{l} = `{r}`" for l, r in zip(labels, a.runs)), "",
         "## 1. 메인 표", "", "저하 시나리오: DR + PreContam + PreDegr + Miss = 1. 비저하: FAR = FAR clean + FAR contam. "
         "무작위 기준선은 매 cycle 독립 확률 p 로 경보 (100회 평균); p 는 각 run 의 cycle 당 경보율.", "", md(main), "",
         "## 2. unit bootstrap 95% CI", "", md(ci), "",
         f"## 3. 기준 run ({ref}) 대비 짝지은 차이", "", "같은 시나리오끼리 차이를 구하고 unit 을 복원추출한 CI. `*` = CI 가 0 을 포함하지 않음.", "", md(diff), "",
         "## 4. 유형별 DR · FAR", "", md(by_type), "",
         "## 5. 센서별 DR · FAR", "", md(by_sensor), "",
         "## 6. 오염 시점별 오염 전 경보율 (저하·비저하 합산)", "", "오염 전 구간이 길수록 (늦은 오염) 오경보가 쌓인다.", "", md(timing), "",
         "run 별 상세: " + " · ".join(f"[{l}]({a.out}_{l}.md)" for l in labels), ""]
    (OUT / f"compare_{a.out}.md").write_text("\n".join(L), encoding="utf-8")

    wb = Workbook(); F, FB = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True); HEAD = PatternFill("solid", fgColor="D9E1F2")
    for title, df in [("메인", main), ("CI", ci), ("짝지은 차이", diff), ("유형별", by_type), ("센서별", by_sensor), ("시점별 오염전 경보", timing)]:
        ws = wb.create_sheet(title[:31])
        for j, c in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=j, value=str(c)); cell.font = FB; cell.fill = HEAD
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                v = None if isinstance(v, float) and pd.isna(v) else (float(v) if hasattr(v, "item") else v)
                cell = ws.cell(row=i, column=j, value=v); cell.font = F
                if isinstance(v, float):
                    cell.number_format = "0.000"
        for j in range(1, len(df.columns) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 16
    wb.remove(wb["Sheet"])
    wb.save(OUT / f"compare_{a.out}.xlsx")
    print(f"→ {OUT / f'compare_{a.out}.md'}  (+ .xlsx, run 별 {a.out}_{{label}}.md)")


if __name__ == "__main__":
    main()
